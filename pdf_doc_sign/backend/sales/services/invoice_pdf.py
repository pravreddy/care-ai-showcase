"""
Invoice PDF Generator — generates professional invoices using reportlab.
Matches Indian business invoice format with GST support.

Respects invoice content options:
  - include_daily_breakdown: show/hide the daily line-item table
  - include_packing_charges: show/hide packing column
  - include_courier_charges: show/hide courier column
  - include_item_detail: show "5x Chapathi @10" vs just food total
"""
import io
import json
from decimal import Decimal

from django.core.files.base import ContentFile

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable,
)
from reportlab.lib.enums import TA_CENTER, TA_RIGHT


def _format_items(items_description, show_detail=True):
    """
    Parse items_description JSON and format as readable text.
    If show_detail=False, just return item count summary.
    """
    if not items_description:
        return '\u2014'
    try:
        items = json.loads(items_description)
        if isinstance(items, list) and items:
            if show_detail:
                parts = []
                for item in items:
                    name = item.get('item_name', item.get('name', '?')).title()
                    qty = item.get('quantity', 1)
                    price = item.get('unit_price', 0)
                    parts.append(f"{qty}x {name} @{price}")
                return ', '.join(parts)
            else:
                return f"{len(items)} item(s)"
    except (json.JSONDecodeError, TypeError, AttributeError):
        pass
    text = str(items_description)
    return text[:60] + '...' if len(text) > 60 else text


def generate_invoice_pdf(invoice):
    """
    Generate a PDF invoice and attach it to the Invoice model.
    Respects invoice.include_* options.
    """
    from sales.models import DailySale

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    # Custom styles
    header_style = ParagraphStyle('InvHeader', parent=styles['Heading2'], fontSize=13, textColor=colors.HexColor('#333333'), spaceBefore=5 * mm)
    normal_style = ParagraphStyle('InvNormal', parent=styles['Normal'], fontSize=9, leading=13)
    bold_style = ParagraphStyle('InvBold', parent=styles['Normal'], fontSize=9, leading=13, fontName='Helvetica-Bold')
    right_style = ParagraphStyle('InvRight', parent=styles['Normal'], fontSize=9, alignment=TA_RIGHT)
    center_style = ParagraphStyle('InvCenter', parent=styles['Normal'], fontSize=9, alignment=TA_CENTER)
    subtitle_style = ParagraphStyle('InvSubtitle', parent=styles['Normal'], fontSize=10, textColor=colors.grey)

    org = invoice.client.organisation
    cs = invoice.currency_symbol
    show_packing = invoice.include_packing_charges
    show_courier = invoice.include_courier_charges
    show_items = invoice.include_item_detail
    show_breakdown = invoice.include_daily_breakdown

    # --- HEADER ---
    header_data = [
        [
            Paragraph(f"<b>{org.name}</b>", ParagraphStyle('OrgName', parent=styles['Normal'], fontSize=16, textColor=colors.HexColor('#1565c0'))),
            Paragraph("<b>INVOICE</b>", ParagraphStyle('InvLabel', parent=styles['Normal'], fontSize=20, alignment=TA_RIGHT, textColor=colors.HexColor('#1565c0'))),
        ],
        [
            Paragraph(f"{org.address or ''}<br/>{org.city or ''} {org.postcode or ''}", normal_style),
            Paragraph(f"<b>Invoice #:</b> {invoice.invoice_number}<br/><b>Date:</b> {invoice.created_at.strftime('%d %b %Y')}", right_style),
        ],
    ]
    if org.company_number:
        header_data.append([Paragraph(f"Company No: {org.company_number}", subtitle_style), ''])
    if org.vat_number:
        header_data.append([Paragraph(f"GST/VAT: {org.vat_number}", subtitle_style), ''])

    header_table = Table(header_data, colWidths=[doc.width * 0.55, doc.width * 0.45])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2 * mm),
    ]))
    elements.append(header_table)
    elements.append(HRFlowable(width='100%', thickness=1, color=colors.HexColor('#1565c0'), spaceAfter=5 * mm))

    # --- BILL TO ---
    client = invoice.client
    elements.append(Paragraph('Bill To:', header_style))
    elements.append(Paragraph(f"<b>{client.invoice_name}</b>", bold_style))
    if client.address:
        elements.append(Paragraph(client.address, normal_style))
    if client.city:
        city_line = client.city
        if client.country:
            city_line += f", {client.country}"
        elements.append(Paragraph(city_line, normal_style))
    if client.contact_person:
        elements.append(Paragraph(f"Attn: {client.contact_person}", normal_style))
    if client.contact_email:
        elements.append(Paragraph(f"Email: {client.contact_email}", normal_style))
    if client.gst_number:
        elements.append(Paragraph(f"GSTIN: {client.gst_number}", normal_style))
    if client.pan_number:
        elements.append(Paragraph(f"PAN: {client.pan_number}", normal_style))

    elements.append(Spacer(1, 3 * mm))
    elements.append(Paragraph(
        f"<b>Period:</b> {invoice.from_date.strftime('%d %b %Y')} to {invoice.to_date.strftime('%d %b %Y')}",
        normal_style,
    ))
    elements.append(Spacer(1, 5 * mm))

    # --- DAILY BREAKDOWN TABLE (optional) ---
    sales = DailySale.objects.filter(
        client=invoice.client,
        date__gte=invoice.from_date,
        date__lte=invoice.to_date,
    ).order_by('date')

    if show_breakdown:
        # Build dynamic columns based on options
        headers = [
            Paragraph('<b>Date</b>', center_style),
            Paragraph('<b>Day</b>', center_style),
        ]
        col_weights = [0.08, 0.06]

        if show_items:
            headers.append(Paragraph('<b>Items</b>', bold_style))
            col_weights.append(0.40)

        headers.append(Paragraph('<b>Food</b>', right_style))
        col_weights.append(0.14 if not show_items else 0.12)

        if show_packing:
            headers.append(Paragraph('<b>Packing</b>', right_style))
            col_weights.append(0.10)
        if show_courier:
            headers.append(Paragraph('<b>Courier</b>', right_style))
            col_weights.append(0.10)

        headers.append(Paragraph('<b>Total</b>', right_style))
        col_weights.append(0.14)

        # Normalize column weights
        total_weight = sum(col_weights)
        col_widths = [doc.width * (w / total_weight) for w in col_weights]

        table_data = [headers]

        for sale in sales:
            if sale.is_holiday:
                row = [
                    Paragraph(sale.date.strftime('%d/%m'), center_style),
                    Paragraph(sale.day_name[:3], center_style),
                ]
                if show_items:
                    row.append(Paragraph(f"<i>{sale.holiday_reason or 'Holiday'}</i>", normal_style))
                row.append(Paragraph('\u2014', right_style))
                if show_packing:
                    row.append(Paragraph('\u2014', right_style))
                if show_courier:
                    row.append(Paragraph('\u2014', right_style))
                row.append(Paragraph(f'{cs}0', right_style))
            else:
                row = [
                    Paragraph(sale.date.strftime('%d/%m'), center_style),
                    Paragraph(sale.day_name[:3], center_style),
                ]
                if show_items:
                    items_text = _format_items(sale.items_description, show_detail=show_items)
                    num_str = f" ({sale.num_people} ppl)" if sale.num_people > 0 else ""
                    row.append(Paragraph(f"{items_text}{num_str}", normal_style))
                row.append(Paragraph(f"{cs}{sale.food_price_per_day:,.0f}", right_style))
                if show_packing:
                    row.append(Paragraph(f"{cs}{sale.packing_charges:,.0f}", right_style))
                if show_courier:
                    row.append(Paragraph(f"{cs}{sale.courier_charges:,.0f}", right_style))
                row.append(Paragraph(f"<b>{cs}{sale.total_charge:,.0f}</b>", right_style))
            table_data.append(row)

        daily_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        daily_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1565c0')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (1, -1), 'CENTER'),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
            ('TOPPADDING', (0, 0), (-1, -1), 2 * mm),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2 * mm),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        elements.append(daily_table)
        elements.append(Spacer(1, 5 * mm))

    # --- TOTALS ---
    tax_label = client.tax_label  # GST / VAT / Sales Tax / Tax
    totals_data = [
        ['Working Days:', f'{invoice.working_days}'],
        ['Total Food:', f'{cs}{invoice.total_food:,.2f}'],
    ]
    if show_packing and invoice.total_packing > 0:
        totals_data.append(['Total Packing:', f'{cs}{invoice.total_packing:,.2f}'])
    if show_courier and invoice.total_courier > 0:
        totals_data.append(['Total Courier:', f'{cs}{invoice.total_courier:,.2f}'])
    if invoice.total_other > 0:
        totals_data.append(['Other Charges:', f'{cs}{invoice.total_other:,.2f}'])
    if invoice.total_discount > 0:
        totals_data.append(['Discount:', f'-{cs}{invoice.total_discount:,.2f}'])

    totals_data.append(['Subtotal:', f'{cs}{invoice.subtotal:,.2f}'])

    if invoice.tax_percentage > 0:
        totals_data.append([f'{tax_label} ({invoice.tax_percentage}%):', f'{cs}{invoice.tax_amount:,.2f}'])

    totals_data.append(['GRAND TOTAL:', f'{cs}{invoice.grand_total:,.2f}'])

    formatted_totals = []
    for i, (label, value) in enumerate(totals_data):
        is_last = i == len(totals_data) - 1
        formatted_totals.append([
            '',
            Paragraph(f"<b>{label}</b>" if is_last else label, right_style),
            Paragraph(f"<b>{value}</b>" if is_last else value, right_style),
        ])

    totals_table = Table(formatted_totals, colWidths=[doc.width * 0.55, doc.width * 0.25, doc.width * 0.20])
    totals_table.setStyle(TableStyle([
        ('LINEABOVE', (1, -1), (-1, -1), 1.5, colors.HexColor('#1565c0')),
        ('TOPPADDING', (0, 0), (-1, -1), 1.5 * mm),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 1.5 * mm),
        ('BACKGROUND', (1, -1), (-1, -1), colors.HexColor('#e8f0fe')),
    ]))
    elements.append(totals_table)

    # --- NOTES ---
    if invoice.notes:
        elements.append(Spacer(1, 5 * mm))
        elements.append(Paragraph('<b>Notes:</b>', bold_style))
        elements.append(Paragraph(invoice.notes, normal_style))

    # --- FOOTER ---
    elements.append(Spacer(1, 10 * mm))
    elements.append(HRFlowable(width='100%', thickness=0.5, color=colors.grey))
    elements.append(Spacer(1, 2 * mm))
    elements.append(Paragraph(
        f"Generated by SignSimple \u2014 {org.name} \u2014 {invoice.created_at.strftime('%d %b %Y %H:%M')}",
        ParagraphStyle('Footer', parent=styles['Normal'], fontSize=7, textColor=colors.grey, alignment=TA_CENTER),
    ))

    # Build PDF
    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()

    filename = f"invoice_{invoice.invoice_number.replace('/', '-')}.pdf"
    invoice.pdf_file.save(filename, ContentFile(pdf_bytes), save=True)

    return pdf_bytes


def generate_invoice_excel(invoice):
    """
    Generate an Excel breakdown of daily sales for attachment.
    Returns (filename, bytes) tuple.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    from sales.models import DailySale

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Daily Breakdown'

    # Styles
    header_font = Font(bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
    bold_font = Font(bold=True, size=10)
    money_fmt = '#,##0.00'
    thin_border = Border(
        bottom=Side(style='thin', color='CCCCCC'),
    )

    org = invoice.client.organisation
    cs = invoice.currency_symbol

    # Title rows
    ws.merge_cells('A1:G1')
    ws['A1'] = f"{org.name} — Invoice {invoice.invoice_number}"
    ws['A1'].font = Font(bold=True, size=14, color='1565C0')

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Client: {invoice.client.invoice_name}"
    ws['A2'].font = Font(bold=True, size=11)

    ws.merge_cells('A3:G3')
    ws['A3'] = f"Period: {invoice.from_date.strftime('%d %b %Y')} to {invoice.to_date.strftime('%d %b %Y')}"

    # Headers (row 5)
    headers = ['Date', 'Day', 'Items', 'Food', 'Packing', 'Courier', 'Total']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    # Data rows
    sales = DailySale.objects.filter(
        client=invoice.client,
        date__gte=invoice.from_date,
        date__lte=invoice.to_date,
    ).order_by('date')

    row = 6
    for sale in sales:
        ws.cell(row=row, column=1, value=sale.date.strftime('%d/%m/%Y'))
        ws.cell(row=row, column=2, value=sale.day_name[:3])

        if sale.is_holiday:
            ws.cell(row=row, column=3, value=sale.holiday_reason or 'Holiday')
            ws.cell(row=row, column=4, value=0).number_format = money_fmt
            ws.cell(row=row, column=5, value=0).number_format = money_fmt
            ws.cell(row=row, column=6, value=0).number_format = money_fmt
            ws.cell(row=row, column=7, value=0).number_format = money_fmt
        else:
            items_text = _format_items(sale.items_description, show_detail=True)
            ws.cell(row=row, column=3, value=items_text)
            ws.cell(row=row, column=4, value=float(sale.food_price_per_day)).number_format = money_fmt
            ws.cell(row=row, column=5, value=float(sale.packing_charges)).number_format = money_fmt
            ws.cell(row=row, column=6, value=float(sale.courier_charges)).number_format = money_fmt
            ws.cell(row=row, column=7, value=float(sale.total_charge)).number_format = money_fmt

        for col in range(1, 8):
            ws.cell(row=row, column=col).border = thin_border
        row += 1

    # Totals row
    row += 1
    ws.cell(row=row, column=3, value='TOTALS').font = bold_font
    ws.cell(row=row, column=4, value=float(invoice.total_food)).number_format = money_fmt
    ws.cell(row=row, column=4).font = bold_font
    ws.cell(row=row, column=5, value=float(invoice.total_packing)).number_format = money_fmt
    ws.cell(row=row, column=5).font = bold_font
    ws.cell(row=row, column=6, value=float(invoice.total_courier)).number_format = money_fmt
    ws.cell(row=row, column=6).font = bold_font
    ws.cell(row=row, column=7, value=float(invoice.subtotal)).number_format = money_fmt
    ws.cell(row=row, column=7).font = bold_font

    if invoice.tax_percentage > 0:
        row += 1
        tax_label = invoice.client.tax_label
        ws.cell(row=row, column=3, value=f'{tax_label} ({invoice.tax_percentage}%)').font = bold_font
        ws.cell(row=row, column=7, value=float(invoice.tax_amount)).number_format = money_fmt

    row += 1
    ws.cell(row=row, column=3, value='GRAND TOTAL').font = Font(bold=True, size=12)
    ws.cell(row=row, column=7, value=float(invoice.grand_total)).number_format = money_fmt
    ws.cell(row=row, column=7).font = Font(bold=True, size=12)

    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 6
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 14

    # Save to bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()

    filename = f"breakdown_{invoice.invoice_number.replace('/', '-')}.xlsx"
    return filename, excel_bytes
