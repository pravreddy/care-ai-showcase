"""
AI Detector — File Text Extraction Utilities
Handles: PDF, DOCX, TXT, PPTX, XLSX
All in one place — easy to extend later.
"""
import io

SUPPORTED_EXTENSIONS = ['pdf', 'docx', 'txt', 'pptx', 'xlsx']


def get_extension(filename: str) -> str:
    return filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''


def extract_text(file, filename: str) -> tuple[str, str]:
    """
    Extract plain text from an uploaded file.
    Returns: (extracted_text, file_type)
    Raises: ValueError if unsupported or unreadable.
    """
    ext = get_extension(filename)

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f'Unsupported file type: .{ext}. Supported: {", ".join(SUPPORTED_EXTENSIONS)}')

    if ext == 'txt':
        return _extract_txt(file), 'txt'
    elif ext == 'pdf':
        return _extract_pdf(file), 'pdf'
    elif ext == 'docx':
        return _extract_docx(file), 'docx'
    elif ext == 'pptx':
        return _extract_pptx(file), 'pptx'
    elif ext == 'xlsx':
        return _extract_xlsx(file), 'xlsx'

    raise ValueError(f'Handler missing for .{ext}')


def _extract_txt(file) -> str:
    raw = file.read()
    for enc in ('utf-8', 'latin-1', 'cp1252'):
        try:
            return raw.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError('Could not decode text file — unsupported encoding.')


def _extract_pdf(file) -> str:
    try:
        import pdfplumber
    except ImportError:
        # Fallback to PyPDF2 which is already in requirements.txt
        return _extract_pdf_pypdf2(file)

    text_parts = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return '\n'.join(text_parts).strip()


def _extract_pdf_pypdf2(file) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(file)
    parts = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            parts.append(t)
    return '\n'.join(parts).strip()


def _extract_docx(file) -> str:
    from docx import Document
    doc = Document(file)
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip()).strip()


def _extract_pptx(file) -> str:
    from pptx import Presentation
    prs = Presentation(file)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                parts.append(shape.text.strip())
    return '\n'.join(parts).strip()


def _extract_xlsx(file) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            row_text = ' '.join(str(c) for c in row if c is not None and str(c).strip())
            if row_text.strip():
                parts.append(row_text)
    return '\n'.join(parts).strip()
